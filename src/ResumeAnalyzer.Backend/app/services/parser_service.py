import io
import docx
from pypdf import PdfReader
import pdfplumber
import logging

logger = logging.getLogger(__name__)

class ParserService:
    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        """
        Extracts plain text from PDF bytes using pdfplumber with a PyPDF fallback.
        """
        extracted_text = ""
        try:
            # Attempt extraction via pdfplumber first (generally cleaner format/layout preservation)
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        extracted_text += text + "\n"
        except Exception as e:
            logger.warning(f"pdfplumber extraction failed: {str(e)}. Falling back to PyPDF.")
            
        # Fallback to PyPDF if pdfplumber returns empty or fails
        if not extracted_text.strip():
            try:
                reader = PdfReader(io.BytesIO(file_bytes))
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        extracted_text += text + "\n"
            except Exception as e:
                logger.error(f"PyPDF extraction also failed: {str(e)}")
                raise ValueError("Failed to extract text from PDF file. File may be corrupted or password-protected.")

        return extracted_text.strip()

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        """
        Extracts plain text from DOCX bytes using python-docx.
        """
        try:
            doc_file = io.BytesIO(file_bytes)
            doc = docx.Document(doc_file)
            
            paragraphs = []
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    paragraphs.append(paragraph.text)
                    
            # Extract from tables too (common in resumes)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.append(cell.text)
                        
            return "\n".join(paragraphs).strip()
        except Exception as e:
            logger.error(f"DOCX text extraction failed: {str(e)}")
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")

    @classmethod
    def extract_text(cls, file_bytes: bytes, filename: str) -> str:
        """
        Determines file type based on extension and extracts text.
        """
        fn = filename.lower()
        if fn.endswith(".pdf"):
            return cls.extract_text_from_pdf(file_bytes)
        elif fn.endswith(".docx"):
            return cls.extract_text_from_docx(file_bytes)
        else:
            raise ValueError("Unsupported file format. Only PDF and DOCX files are allowed.")
